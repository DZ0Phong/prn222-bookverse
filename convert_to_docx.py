import os
try:
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    print("python-docx not installed.")
    exit(1)

md_file = "Project_Specifications.md"
docx_file = "Project_Specifications.docx"

if not os.path.exists(md_file):
    print("Markdown file not found!")
    exit(1)

document = Document()

# Add a title
title = document.add_heading('Đặc tả dự án Bookverse', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

with open(md_file, "r", encoding="utf-8") as f:
    for line in f:
        line = line.strip()
        if not line:
            continue
        
        if line.startswith("# "):
            document.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=2)
        elif line.startswith("### "):
            document.add_heading(line[4:], level=3)
        elif line.startswith("- **") or line.startswith("**"):
            p = document.add_paragraph()
            p.style = 'List Bullet'
            p.add_run(line.replace("**", "").replace("- ", ""))
        elif line.startswith("- "):
            p = document.add_paragraph(line[2:], style='List Bullet')
        else:
            document.add_paragraph(line)

document.save(docx_file)
print(f"Successfully generated {docx_file}")
