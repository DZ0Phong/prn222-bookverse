from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "documents"
OUT.mkdir(parents=True, exist_ok=True)
TARGET = OUT / "BookVerse_Requirement_Design_Specification.docx"

BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
PALE = "F3F6F9"
ORANGE = "F4B183"
GRAY = "666666"
WHITE = "FFFFFF"
BLACK = "000000"


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_width(cell, dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[min(idx, len(widths) - 1)])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def format_run(run, size=10.5, bold=False, color=BLACK, italic=False, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    format_run(run, 9, color=GRAY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.18
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        format_run(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        format_run(r2)
    else:
        format_run(p.add_run(text))
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        format_run(p.add_run(item))


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        format_run(p.add_run(item))


def add_table(doc, headers, rows, widths=None, header_fill=LIGHT_BLUE, font_size=9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, header in enumerate(headers):
        set_cell_fill(hdr.cells[i], header_fill)
        p = hdr.cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        format_run(p.add_run(str(header)), font_size, bold=True, color=BLUE)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            p = cells[i].paragraphs[0]
            format_run(p.add_run(str(value)), font_size)
    set_table_geometry(table, widths or [int(9360 / len(headers))] * len(headers))
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_code(doc, code):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_fill(cell, "F7F7F7")
    set_cell_margins(cell, 120, 160, 120, 160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    for idx, line in enumerate(code.strip().splitlines()):
        if idx:
            p.add_run("\n")
        format_run(p.add_run(line), 8.3, name="Consolas")
    set_table_geometry(table, [9360])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_flow(doc, nodes):
    table = doc.add_table(rows=1, cols=len(nodes))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    widths = [int(9360 / len(nodes))] * len(nodes)
    for i, node in enumerate(nodes):
        cell = table.cell(0, i)
        set_cell_fill(cell, LIGHT_BLUE if i % 2 == 0 else "E2F0D9")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        format_run(p.add_run(node), 8.5, bold=True, color=BLUE)
    set_table_geometry(table, widths)
    p = doc.add_paragraph("  →  ".join(nodes))
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    for r in p.runs:
        format_run(r, 8, color=GRAY)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18
    for level, size, before, after in ((1, 16, 16, 8), (2, 13, 12, 6), (3, 11.5, 8, 4)):
        style = styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLUE)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)
section.header_distance = Inches(0.3)
section.footer_distance = Inches(0.3)
configure_styles(doc)

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
format_run(header.add_run("PRN222 | BOOKVERSE"), 8.5, bold=True, color=GRAY)
add_page_number(section.footer.paragraphs[0])

# Cover
for _ in range(4):
    doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
format_run(p.add_run("FPT UNIVERSITY"), 13, bold=True, color=ORANGE)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(28)
format_run(p.add_run("REQUIREMENT & DESIGN\nSPECIFICATION"), 24, bold=True, color=BLACK)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(10)
format_run(p.add_run("BOOKVERSE - ONLINE BOOKSTORE MANAGEMENT SYSTEM"), 15, bold=True, color=BLUE)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(26)
format_run(p.add_run("Subject: PRN222"), 11, bold=True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
format_run(p.add_run("Group: ____________________"), 11)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
format_run(p.add_run("Semester: Summer 2026"), 11)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(90)
format_run(p.add_run("Hanoi, 2026"), 10, italic=True, color=GRAY)
doc.add_page_break()

# Contents
add_heading(doc, "Contents", 1)
contents = [
    ("Record of Changes", "3"),
    ("Project Plan and Work Allocation", "4"),
    ("I. Overview", "5"),
    ("II. User Requirements", "6"),
    ("III. Project Objectives", "10"),
    ("IV. Database Design", "11"),
    ("V. Main Functions of Application", "17"),
    ("VI. Project Structure and Implementation", "20"),
    ("VII. Test Cases", "24"),
    ("Appendix A. Database Setup", "26"),
]
add_table(doc, ["Section", "Page"], contents, [8200, 1160], header_fill=ORANGE, font_size=9.5)
add_body(doc, "The table of contents follows the structure of the school reference report. Page numbers may be refreshed in Microsoft Word after final team edits.")
doc.add_page_break()

add_heading(doc, "Record of Changes", 1)
add_table(
    doc,
    ["Version", "Date", "A/M/D", "In charge", "Change Description"],
    [
        ("v0.1", "____/____/2026", "A", "________________", "Initial requirement and design specification."),
        ("v0.2", "____/____/2026", "M", "________________", "Update database and application workflows."),
        ("v1.0", "____/____/2026", "A", "________________", "Final review and submission."),
    ],
    [950, 1500, 900, 1900, 4110],
    header_fill=ORANGE,
)
add_body(doc, "Legend: A - Added, M - Modified, D - Deleted.")

add_heading(doc, "Project Plan and Work Allocation", 1)
add_table(
    doc,
    ["No.", "Work package", "Member", "Main responsibility", "Status"],
    [
        ("1", "Authentication, User & Product", "________________", "Login/register, authorization, user administration, books, categories, inventory, search.", "________"),
        ("2", "User Commerce & Review", "________________", "Cart, checkout, voucher, order history, order tracking, review and rating.", "________"),
        ("3", "Admin Operation & Reporting", "________________", "Dashboard, order processing, review moderation, revenue reports, responsive UI.", "________"),
    ],
    [600, 1850, 1750, 4360, 800],
)
add_body(doc, "Member names and student IDs are intentionally left blank for the team to complete.")
doc.add_page_break()

# Overview
add_heading(doc, "I. Overview", 1)
add_body(doc, "BookVerse is an ASP.NET Core MVC web application for online bookstore management. The system supports customers in discovering books, maintaining accounts, viewing purchase history, and submitting book reviews. Administrators manage catalog data, users, vouchers, orders, review moderation, dashboards, and business reports.")
add_heading(doc, "1.1 Context Diagram", 2)
add_flow(doc, ["Guest / User", "BookVerse MVC Application", "SQL Server Database", "Administrator"])
add_bullets(doc, [
    "Guests browse the catalog and register or sign in.",
    "Authenticated users access account-specific functions, order history, and reviews.",
    "Administrators operate catalog, user, order, voucher, review, dashboard, and reporting modules.",
    "Entity Framework Core maps C# entities and LINQ queries to the QuanLyBanSach SQL Server database.",
])
add_heading(doc, "1.2 Technology Stack", 2)
add_table(doc, ["Layer", "Technology", "Purpose"], [
    ("Presentation", "Razor Views, HTML, CSS, Bootstrap/JavaScript", "Render pages and receive user interactions."),
    ("Application", "ASP.NET Core MVC (.NET 8)", "Routing, controllers, authentication, authorization, and workflow processing."),
    ("Data access", "Entity Framework Core 8", "Map entities, execute LINQ queries, and persist changes."),
    ("Database", "Microsoft SQL Server", "Store users, catalog, cart, orders, vouchers, reviews, and wishlist."),
    ("Supporting", "BCrypt.Net, ClosedXML", "Password hashing and Excel report export."),
], [1600, 2400, 5360])

# Requirements
add_heading(doc, "II. User Requirements", 1)
add_heading(doc, "2.1 Actors", 2)
add_table(doc, ["No.", "Actor", "Description"], [
    ("1", "Guest", "Unauthenticated visitor who can browse books, search, register, and sign in."),
    ("2", "User / Customer", "Authenticated customer who uses account-specific shopping and review functions."),
    ("3", "Administrator", "Authorized operator who manages users, catalog, orders, vouchers, reviews, dashboard, and reports."),
], [600, 1900, 6860])

add_heading(doc, "2.2 Use Case Overview", 2)
add_table(doc, ["Actor", "Use cases"], [
    ("Guest", "Browse home page; view book detail; search books; register; log in."),
    ("User", "Log out; view order history; follow order/payment status; submit rating and comment; use planned cart and checkout functions."),
    ("Administrator", "Manage users; books; categories; orders; vouchers; reviews; dashboard; statistics; export Excel reports."),
], [1800, 7560])

add_heading(doc, "2.3 Use Case Descriptions", 2)
use_cases = [
    ("UC-01", "Register", "Guest", "Create a customer account with validated email and BCrypt password."),
    ("UC-02", "Login", "Guest/User/Admin", "Verify credentials, create cookie claims, and redirect by role."),
    ("UC-03", "View book detail", "Guest/User", "Load a book, category, and approved reviews."),
    ("UC-04", "Search and browse", "Guest/User", "Filter active books and display catalog sections."),
    ("UC-05", "View order history", "User", "Load orders belonging to the authenticated user and show status."),
    ("UC-06", "Submit review", "User", "Create a rating/comment in pending approval status."),
    ("UC-07", "Manage users", "Admin", "Search users and lock/unlock non-admin accounts."),
    ("UC-08", "Manage catalog", "Admin", "Create, update, hide/delete books and manage categories."),
    ("UC-09", "Manage orders", "Admin", "Filter orders, inspect detail, and update valid status transitions."),
    ("UC-10", "Manage vouchers", "Admin", "Create, edit, activate, deactivate, or delete discount codes."),
    ("UC-11", "Moderate reviews", "Admin", "Approve pending reviews or remove inappropriate reviews."),
    ("UC-12", "Dashboard and reports", "Admin", "Review KPIs and export order/revenue reports to Excel."),
]
add_table(doc, ["ID", "Use Case", "Actor", "Description"], use_cases, [900, 1900, 1700, 4860])

add_heading(doc, "2.4 Detailed Workflow - Order History", 2)
add_flow(doc, ["GET /Account/OrderHistory", "[Authorize]", "Read UserId claim", "Query Orders", "Map ViewModel", "Render View"])
add_numbered(doc, [
    "The browser requests /Account/OrderHistory.",
    "Cookie authentication checks [Authorize]. Unauthenticated requests redirect to /Account/Login.",
    "AccountController reads ClaimTypes.NameIdentifier and parses the current UserId.",
    "Entity Framework filters Orders by UserId, orders them by CreatedAt, and sums OrderDetail quantities.",
    "The query maps results to List<UserOrderListViewModel>.",
    "OrderHistory.cshtml receives the list; an empty list shows the empty-state message, otherwise Razor renders each order row.",
])

add_heading(doc, "2.5 Detailed Workflow - Book Review", 2)
add_flow(doc, ["GET Book Detail", "Load approved reviews", "Render form", "POST Review/Create", "Save pending review", "Admin approval"])
add_numbered(doc, [
    "BookController.Detail receives the book ID from the route and returns HTTP 404 when the book does not exist.",
    "The controller loads reviews where BookId matches and IsApproved is true, then packages Book and ExistingReviews in BookDetailViewModel.",
    "An authenticated user submits BookId, Rating, Comment, and an anti-forgery token to /Review/Create.",
    "ReviewController obtains UserId from the authentication claim and creates Review with IsApproved = false.",
    "SaveChangesAsync inserts the review. TempData stores a success message and HTTP 302 redirects to the book detail page.",
    "The review becomes visible after an administrator changes IsApproved to true.",
])

# Objectives
add_heading(doc, "III. Project Objectives", 1)
add_bullets(doc, [
    "Apply ASP.NET Core MVC and Razor Views to a practical bookstore domain.",
    "Implement cookie-based authentication, claims, and role-based authorization.",
    "Use Entity Framework Core with SQL Server following a database-first model.",
    "Provide catalog and administrative CRUD workflows with validation.",
    "Support order visibility, review moderation, voucher administration, dashboards, and reports.",
    "Maintain a responsive interface and clear separation between Models, Views, Controllers, and ViewModels.",
])
add_heading(doc, "3.1 Current Scope Status", 2)
add_table(doc, ["Module", "Current implementation status"], [
    ("Authentication and authorization", "Implemented: register, login, logout, cookie claims, Admin/User role checks."),
    ("Catalog administration", "Implemented: books, categories, inventory fields, image handling, filters, pagination."),
    ("User order history", "Implemented: list and status display. User order detail page is not yet implemented."),
    ("Review", "Implemented: submit pending review, list approved reviews, admin approve/delete."),
    ("Cart and checkout", "Database entities exist, but user-side controller/view/action workflow is not connected."),
    ("Voucher", "Admin CRUD implemented; user-side apply-voucher workflow is not connected."),
    ("Admin order/reporting", "Implemented: order filtering/detail/status changes, dashboard, Excel report."),
], [2700, 6660], header_fill=ORANGE)

# Database
add_heading(doc, "IV. Database Design", 1)
add_heading(doc, "4.1 ERD - Relationship Summary", 2)
add_table(doc, ["Parent entity", "Relationship", "Child entity", "Foreign key"], [
    ("Roles", "1 - N", "Users", "Users.RoleId"),
    ("Users", "1 - N", "Cart", "Cart.UserId"),
    ("Books", "1 - N", "Cart", "Cart.BookId"),
    ("Users", "1 - N", "Orders", "Orders.UserId"),
    ("Vouchers", "1 - N", "Orders", "Orders.VoucherId"),
    ("Orders", "1 - N", "OrderDetails", "OrderDetails.OrderId"),
    ("Books", "1 - N", "OrderDetails", "OrderDetails.BookId"),
    ("Users", "1 - N", "Reviews", "Reviews.UserId"),
    ("Books", "1 - N", "Reviews", "Reviews.BookId"),
    ("Users", "1 - N", "Wishlist", "Wishlist.UserId"),
    ("Books", "1 - N", "Wishlist", "Wishlist.BookId"),
    ("Categories", "1 - N", "Books", "Books.CategoryId"),
], [2050, 1200, 2400, 3710])

add_heading(doc, "4.2 Entity Overview", 2)
entities = {
    "Roles": "RoleId (PK), RoleName",
    "Users": "UserId (PK), FullName, Email, Password, Phone, Address, RoleId (FK), IsActive, CreatedAt",
    "Categories": "CategoryId (PK), CategoryName, Description",
    "Books": "BookId (PK), Title, Author, Publisher, CategoryId (FK), Price, Quantity, Image, Description, IsActive, CreatedAt",
    "Cart": "CartId (PK), UserId (FK), BookId (FK), Quantity",
    "Vouchers": "VoucherId (PK), Code, DiscountPercent, ExpiryDate, IsActive",
    "Orders": "OrderId (PK), UserId (FK), VoucherId (FK), TotalAmount, ShippingAddress, PhoneNumber, Status, PaymentMethod, PaymentStatus, CreatedAt",
    "OrderDetails": "OrderDetailId (PK), OrderId (FK), BookId (FK), Quantity, Price",
    "Reviews": "ReviewId (PK), BookId (FK), UserId (FK), Rating, Comment, IsApproved, CreatedAt",
    "Wishlist": "WishlistId (PK), UserId (FK), BookId (FK)",
}
add_table(doc, ["Entity", "Attributes"], list(entities.items()), [1900, 7460])

add_heading(doc, "4.3 Important Data Rules", 2)
add_bullets(doc, [
    "User Email, RoleName, CategoryName, Voucher Code, Cart(UserId, BookId), and Wishlist(UserId, BookId) are unique in the portable setup script.",
    "Book price and quantity cannot be negative.",
    "Cart and OrderDetail quantity must be greater than zero.",
    "Review rating must be between 1 and 5.",
    "Order statuses used by the application are: Chờ xử lý, Đang giao, Đã giao, and Đã hủy.",
    "Pending reviews use IsApproved = false; public book detail only reads approved reviews.",
])

add_heading(doc, "4.4 Database Connection", 2)
add_code(doc, '''
"ConnectionStrings": {
  "MyCnn": "Data Source=.;Initial Catalog=QuanLyBanSach;
            Trusted_Connection=SSPI;Encrypt=false;
            TrustServerCertificate=true"
}''')
add_body(doc, "Program.cs registers QuanLyBanSachContext with UseSqlServer. The corrected db.sql creates the database without machine-specific MDF/LDF paths and safely seeds demo records.")

# Functions
add_heading(doc, "V. Main Functions of Application", 1)
add_heading(doc, "5.1 Guest", 2)
add_bullets(doc, [
    "Read: browse the home page and book details.",
    "Create: register a new user account.",
    "Authentication: submit credentials and receive an eight-hour authentication cookie.",
])
add_heading(doc, "5.2 User / Customer", 2)
add_bullets(doc, [
    "Read personal order history and order/payment status.",
    "Read book detail and approved reviews.",
    "Create a book rating and comment in pending state.",
    "Planned but not connected in the current source: cart add/update/delete, checkout, simulated payment, and voucher application.",
])
add_heading(doc, "5.3 Administrator", 2)
add_bullets(doc, [
    "Books: list, filter, create, edit, upload image, hide/delete where safe.",
    "Categories: create, update, delete with dependency validation.",
    "Users: search and lock/unlock customer accounts.",
    "Orders: filter, inspect detail, and transition status.",
    "Vouchers: create, edit, activate/deactivate, and delete or disable used vouchers.",
    "Reviews: filter, approve, and delete.",
    "Dashboard and reports: KPIs, revenue trends, top books, category revenue, and Excel export.",
])

add_heading(doc, "5.4 Order Status Processing", 2)
add_table(doc, ["Current status", "Allowed next status", "Additional processing"], [
    ("Chờ xử lý", "Đang giao; Đã hủy", "Cancel changes payment status to Đã hủy."),
    ("Đang giao", "Đã giao; Đã hủy", "COD delivery changes payment status to Đã thanh toán."),
    ("Đã giao", "None", "Final state."),
    ("Đã hủy", "None", "Final state."),
], [2000, 2500, 4860])

# Structure/Implementation
add_heading(doc, "VI. Project Structure and Implementation", 1)
add_heading(doc, "6.1 ASP.NET Core MVC Structure", 2)
add_flow(doc, ["Browser", "Routing", "Controller Action", "EF Core DbContext", "SQL Server", "Razor View"])
add_table(doc, ["Folder / File", "Responsibility"], [
    ("Program.cs", "Registers MVC, DbContext, cookie authentication, authorization, session, middleware, and routes."),
    ("Controllers/", "Receives HTTP requests, validates input, runs application workflows, and returns results."),
    ("Controllers/Admin/", "Admin-only operational modules."),
    ("Models/", "Database entities scaffolded for Entity Framework Core."),
    ("Models/ViewModels/", "Shapes data specifically for forms, lists, dashboards, detail pages, and reports."),
    ("Views/", "Razor pages that render HTML from controller-provided models."),
    ("wwwroot/", "Static CSS, JavaScript, icons, and client libraries."),
    ("appsettings.json", "Logging, host configuration, and SQL Server connection string."),
], [2600, 6760])

add_heading(doc, "6.2 Startup Pipeline", 2)
add_code(doc, '''
builder.Services.AddControllersWithViews();
builder.Services.AddDbContext<QuanLyBanSachContext>(options =>
    options.UseSqlServer(builder.Configuration.GetConnectionString("MyCnn")));
builder.Services.AddAuthentication(CookieAuthenticationDefaults.AuthenticationScheme)
    .AddCookie(options => {
        options.LoginPath = "/Account/Login";
        options.AccessDeniedPath = "/Account/AccessDenied";
    });

app.UseRouting();
app.UseSession();
app.UseAuthentication();
app.UseAuthorization();
app.MapControllerRoute(
    name: "default",
    pattern: "{controller=Home}/{action=HomePage}/{id?}");''')

add_heading(doc, "6.3 Order History Mapping", 2)
add_code(doc, '''
var orders = await _context.Orders
    .Where(o => o.UserId == userId)
    .OrderByDescending(o => o.CreatedAt)
    .Select(o => new UserOrderListViewModel
    {
        OrderId = o.OrderId,
        TotalItems = o.OrderDetails.Sum(d => d.Quantity ?? 0),
        TotalAmount = o.TotalAmount ?? 0,
        Status = o.Status,
        PaymentStatus = o.PaymentStatus
    })
    .ToListAsync();

return View(orders);''')

add_heading(doc, "6.4 Review Creation Mapping", 2)
add_code(doc, '''
var newReview = new Review
{
    BookId = BookId,
    UserId = userId,
    Rating = Rating,
    Comment = Comment.Trim(),
    IsApproved = false,
    CreatedAt = DateTime.Now
};

_context.Reviews.Add(newReview);
await _context.SaveChangesAsync();
return RedirectToAction("Detail", "Book", new { id = BookId });''')

add_heading(doc, "6.5 Known Gaps and Risks", 2)
add_table(doc, ["Priority", "Issue", "Recommended action"], [
    ("High", "Cart button only displays a JavaScript alert.", "Implement CartController, cart ViewModels, views, server-side stock validation, and anti-forgery protection."),
    ("High", "No user checkout/order creation transaction.", "Create Order + OrderDetails, reduce stock, clear cart, and commit in one database transaction."),
    ("Medium", "User voucher application is missing.", "Validate code, active flag, expiry, and calculate discount on the server."),
    ("Medium", "Review action does not validate rating range or book existence.", "Validate Rating 1-5 and confirm BookId before insertion."),
    ("Medium", "No user order detail action/view.", "Use the existing UserOrderDetailViewModel and add ownership validation."),
    ("Low", "Some UI names and cart counts are hard-coded.", "Bind authenticated identity and calculated counts."),
], [900, 3540, 4920], header_fill=ORANGE)

# Tests
add_heading(doc, "VII. Test Cases", 1)
test_rows = [
    ("TC-AUTH-01", "Login with active User account", "Valid email/password", "Cookie created; redirect to home."),
    ("TC-AUTH-02", "Login with Admin account", "Valid admin credentials", "Cookie contains Admin role; redirect dashboard."),
    ("TC-AUTH-03", "Login with locked account", "IsActive = false", "Login rejected with locked message."),
    ("TC-BOOK-01", "Open existing book", "Valid BookId", "HTTP 200 and book detail rendered."),
    ("TC-BOOK-02", "Open missing book", "Unknown BookId", "HTTP 404."),
    ("TC-ORDER-01", "View own order history", "Authenticated demo user", "Only matching UserId orders displayed."),
    ("TC-ORDER-02", "View history without login", "Anonymous request", "Redirect to Login."),
    ("TC-REV-01", "Submit valid review", "Rating 1-5 and comment", "Review inserted with IsApproved=false."),
    ("TC-REV-02", "View pending review", "IsApproved=false", "Review not displayed publicly."),
    ("TC-REV-03", "Admin approves review", "Pending ReviewId", "IsApproved=true and review becomes visible."),
    ("TC-ADMIN-01", "Update order to delivered", "COD order in Đang giao", "Status=Đã giao; PaymentStatus=Đã thanh toán."),
    ("TC-ADMIN-02", "Update final order", "Order already Đã giao", "Request rejected."),
    ("TC-DB-01", "Execute db.sql on empty instance", "SQL Server permission to create DB", "Database, schema, indexes, and demo rows created."),
    ("TC-DB-02", "Execute db.sql again", "Existing database", "No duplicate seed records."),
]
add_table(doc, ["ID", "Test objective", "Input / precondition", "Expected result"], test_rows, [1200, 2600, 2600, 2960])
add_heading(doc, "7.1 Build Verification", 2)
add_body(doc, "The current solution builds successfully with 0 errors. Existing warnings concern scaffolded connection-string guidance and nullable reference analysis in ReviewController and BookDetailViewModel.")

# Appendix
add_heading(doc, "Appendix A. Database Setup", 1)
add_heading(doc, "A.1 Corrected Script", 2)
add_body(doc, "Run db.sql from the repository root in SQL Server Management Studio. The script creates QuanLyBanSach without fixed file paths, creates missing tables and indexes, and inserts idempotent demo data.")
add_heading(doc, "A.2 Demo Accounts", 2)
add_table(doc, ["Role", "Email", "Password", "Purpose"], [
    ("Admin", "admin@bookverse.vn", "Admin@123", "Dashboard and administration demonstration."),
    ("User", "user@bookverse.vn", "User@123", "Order history, cart seed, wishlist, and review demonstration."),
    ("User", "lan@bookverse.vn", "User@123", "Additional order and review records."),
], [1500, 3000, 1800, 3060], header_fill=ORANGE)
add_body(doc, "Security note: demo seed passwords are plain text only because the current Login action explicitly supports a legacy plain-text fallback. Production seed data should use BCrypt hashes and the fallback should be removed.")
add_heading(doc, "A.3 Local Connection Troubleshooting", 2)
add_numbered(doc, [
    "Verify SQL Server (MSSQLSERVER or SQLEXPRESS) is running.",
    "In SSMS, test Windows Authentication first. If SSPI fails, repair the Windows/SQL Server service account or use a valid SQL login.",
    "Match appsettings.json Data Source to the real instance, for example '.', '.\\SQLEXPRESS', or '(localdb)\\MSSQLLocalDB'.",
    "Keep Encrypt=false or TrustServerCertificate=true for local development only.",
    "Execute db.sql, verify the summary counts, then run the BookVerse project.",
])
add_heading(doc, "Appendix B. Diagram Maintenance", 1)
add_body(doc, "The report uses Word-native flow and relationship tables so it remains editable without an external service. If the school requires formal UML images, recreate the same content using diagrams.net (draw.io), PlantUML, Mermaid, or Visual Studio class diagrams, export as SVG/PNG, and replace the corresponding Word-native figure.")
add_table(doc, ["Diagram", "Recommended tool", "Source of truth"], [
    ("Context diagram", "diagrams.net or Mermaid", "Actors and application boundaries in Section I."),
    ("Use case diagrams", "diagrams.net or PlantUML", "Actor/use-case list in Section II."),
    ("ERD", "SSMS Database Diagrams or diagrams.net", "db.sql and QuanLyBanSachContext.cs."),
    ("MVC workflow", "Mermaid or diagrams.net", "Program.cs, controllers, ViewModels, and views."),
], [2200, 2500, 4660])

doc.core_properties.title = "BookVerse Requirement & Design Specification"
doc.core_properties.subject = "PRN222"
doc.core_properties.author = "BookVerse Team"
doc.core_properties.keywords = "BookVerse, PRN222, ASP.NET Core MVC, SQL Server"
doc.save(TARGET)
print(TARGET)
